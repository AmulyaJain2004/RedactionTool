"""The upload page: valid documents come back redacted, bad uploads get a clear error."""
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import docx

try:
    from app import app
except ImportError:                                  # flask is optional (requirements-web.txt)
    app = None


@unittest.skipIf(app is None, "flask is not installed")
class WebApp(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def upload(self, data, name):
        return self.client.post("/redact", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")

    def test_pages(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        self.assertEqual(self.client.get("/health").get_json(), {"status": "ok"})

    def test_valid_document_is_redacted(self):
        document = docx.Document()
        document.add_paragraph("Mail rohan.dey@gmail.com or call +91 98765 43210")
        buffer = io.BytesIO()
        document.save(buffer)
        response = self.upload(buffer.getvalue(), "My Doc.docx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("My Doc_redacted.docx", response.headers["Content-Disposition"])
        text = docx.Document(io.BytesIO(response.data)).paragraphs[0].text
        self.assertNotIn("rohan.dey@gmail.com", text)
        self.assertNotIn("98765", text)

    def test_bad_uploads_are_rejected(self):
        self.assertEqual(self.client.post("/redact", data={}, content_type="multipart/form-data").status_code, 400)
        self.assertEqual(self.upload(b"hello", "notes.txt").status_code, 400)
        self.assertEqual(self.upload(b"not a zip", "fake.docx").status_code, 400)


if __name__ == "__main__":
    unittest.main()
