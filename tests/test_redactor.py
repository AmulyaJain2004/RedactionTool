"""Unit tests. Run from the project folder:  python -m unittest discover -s tests -v"""
import os
import re
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from pii_redactor import Redactor
from pii_redactor.fakes import FakeFactory, luhn_check_digit


def redact(text: str, *context: str) -> str:
    """Redact `text`; `context` paragraphs are only used for learning (names, companies)."""
    redactor = Redactor()
    redactor.learn([text, *context])
    return redactor.redact_text(text)


def spans(text: str, *context: str):
    redactor = Redactor()
    redactor.learn([text, *context])
    return [(s.type, s.text) for s in redactor.detect(text)]


class StructuredIdentifiers(unittest.TestCase):
    def test_email(self):
        self.assertEqual(spans("mail rashhi.patil@gmail.com now"), [("EMAIL", "rashhi.patil@gmail.com")])
        self.assertTrue(redact("rohan.dey@gmail.com").endswith("@example.com"))

    def test_phone_formats(self):
        for number in ["+91 9876543210", "+ 91 20 4505 3237", "022-68052182", "(555) 123-4567", "+1 415 555 0132"]:
            self.assertEqual(spans(f"Call {number} today"), [("PHONE", number)], number)

    def test_phone_keeps_country_code_and_format(self):
        out = redact("+91 9876543210")
        self.assertRegex(out, r"^\+91 \d{10}$")
        self.assertNotEqual(out, "+91 9876543210")

    def test_ssn(self):
        self.assertEqual(spans("SSN 123-45-6789 on file"), [("SSN", "123-45-6789")])
        self.assertEqual(spans("Social Security Number: 123456789"), [("SSN", "123456789")])
        self.assertEqual(spans("ref 000-12-3456"), [])                       # impossible area numbers: 000, 666, 9xx
        self.assertEqual(spans("ref 666-12-3456"), [])
        self.assertEqual(spans("ref 912-12-3456"), [])

    def test_credit_card_requires_luhn(self):
        self.assertEqual(spans("card 4111 1111 1111 1111"), [("CREDIT_CARD", "4111 1111 1111 1111")])
        self.assertEqual(spans("order 4111 1111 1111 1112"), [])            # fails Luhn: an order number
        fake = redact("4111-1111-1111-1111")
        digits = re.sub(r"\D", "", fake)
        self.assertEqual(luhn_check_digit(digits[:-1]), digits[-1])   # fake is Luhn-valid too

    def test_ip_addresses(self):
        self.assertEqual(spans("from 10.20.30.40 to"), [("IP_ADDRESS", "10.20.30.40")])
        self.assertEqual(spans("addr 2001:0db8:85a3:0000:0000:8a2e:0370:7334"),
                         [("IP_ADDRESS", "2001:0db8:85a3:0000:0000:8a2e:0370:7334")])
        self.assertEqual(spans("version 1.2.3.4.5 and 999.1.1.1 at 10:30:45"), [])

    def test_date_of_birth_needs_a_cue(self):
        self.assertEqual(spans("DOB: 12/03/1985"), [("DOB", "12/03/1985")])
        self.assertEqual(spans("born on March 4, 1990"), [("DOB", "March 4, 1990")])
        self.assertEqual(spans("Date of Birth - 4th March 1990"), [("DOB", "4th March 1990")])
        self.assertEqual(spans("Dated December 10, 2025"), [])              # ordinary date

    def test_precision_identifiers_are_not_pii(self):
        text = "Order 1234567890 Ticket #98765 CIN U28129PN1979PLC141032 SEBI INM000013004 pin 141032"
        self.assertEqual(spans(text), [])


class Addresses(unittest.TestCase):
    def test_indian_address_with_state_and_country(self):
        text = "Registered Office: 11/3, Village Birdewadi, Chakan Taluka - Khed, Pune - 410 501, Maharashtra, India;"
        found = spans(text)
        expected = "11/3, Village Birdewadi, Chakan Taluka - Khed, Pune - 410 501, Maharashtra, India"
        self.assertEqual(found, [("ADDRESS", expected)])

    def test_us_address(self):
        self.assertEqual(spans("Ship to 742 Evergreen Terrace Street, Springfield, IL 62704 please")[0][0], "ADDRESS")

    def test_prose_before_address_is_kept(self):
        text = "The office is situated at 201, Tower 2, Baner, Pune - 411 045, India"
        (kind, value), = spans(text)
        self.assertEqual(kind, "ADDRESS")
        self.assertTrue(value.startswith("201"))

    def test_registration_number_is_not_an_address(self):
        self.assertEqual(spans("Registration number: 141032"), [])

    def test_address_split_over_lines_of_one_block(self):
        redactor = Redactor()
        block = "11/3, Village Birdewadi Chakan\nPune - 410 501\nMaharashtra, India"
        redactor.learn([block])
        (span,) = redactor.detect(block)
        self.assertEqual((span.start, span.end), (0, len(block)))
        # a heading line above the address must not be swallowed
        block = "REGISTERED OFFICE\n11/3, Village Birdewadi Chakan\nPune - 410 501"
        (span,) = redactor.detect(block)
        self.assertTrue(span.text.startswith("11/3"))


class NamesAndCompanies(unittest.TestCase):
    def test_spec_example_names(self):
        out = redact("Rashi Patil and Rohan Dey")
        self.assertNotIn("Rashi", out)
        self.assertNotIn("Rohan", out)

    def test_same_person_gets_same_fake_in_every_form(self):
        redactor = Redactor()
        redactor.learn(["Kushal Subbayya Hegde",
                        "our Promoter, Chairman and Executive Director, Kushal Subbayya Hegde"])
        a = redactor.redact_text("Kushal Subbayya Hegde")
        b = redactor.redact_text("KUSHAL SUBBAYYA HEGDE")
        c = redactor.redact_text("KushalSubbayya Hegde")
        self.assertEqual(a.upper(), b)
        self.assertEqual(a, c)
        self.assertNotIn("Hegde", a)

    def test_contact_person_list(self):
        found = spans("Contact Person: Eric Bacha/ Sachin Gawade/ Pravin Teli")
        self.assertEqual([t for _, t in found], ["Eric Bacha", "Sachin Gawade", "Pravin Teli"])

    def test_glossary_terms_are_not_names(self):
        for term in ["Indian Rupees", "Basic Custom Duty", "Low Tension", "Book Building Process"]:
            self.assertEqual(spans(term, "the book building process is used"), [], term)

    def test_company_names(self):
        found = spans("Audited by Kirtane & Pandit LLP and banked with HDFC Bank Limited.")
        self.assertEqual([t for _, t in found], ["Kirtane & Pandit LLP", "HDFC Bank Limited"])

    def test_company_short_forms_share_a_fake(self):
        redactor = Redactor()
        texts = ["KSH International Limited", "KSH Infra Park Private Limited", "KSH is a group"]
        redactor.learn(texts)
        full = redactor.redact_text("KSH International Limited")
        brand = redactor.redact_text("KSH is a group")
        self.assertTrue(brand.lower().startswith(full.split()[0].lower()))
        self.assertNotIn("KSH", full + brand)

    def test_public_bodies_are_kept(self):
        text = "listed on the BSE Limited and National Stock Exchange of India Limited under SEBI"
        self.assertEqual(spans(text), [])

    def test_uppercase_suffix(self):
        self.assertEqual([t for _, t in spans("WATERLOO INDUSTRIAL PARK VI PRIVATE LIMITED")],
                         ["WATERLOO INDUSTRIAL PARK VI PRIVATE LIMITED"])


class Consistency(unittest.TestCase):
    def test_same_value_same_fake_different_values_different_fakes(self):
        redactor = Redactor()
        text = "a@x.com b@x.com a@x.com"
        redactor.learn([text])
        first, second, third = redactor.redact_text(text).split()
        self.assertEqual(first, third)
        self.assertNotEqual(first, second)

    def test_fakes_never_equal_the_original(self):
        factory = FakeFactory()
        for kind, value in [("phone", "+91 9876543210"), ("ssn", "123-45-6789"), ("ip_address", "10.0.0.1"),
                            ("credit_card", "4111 1111 1111 1111"), ("dob", "12/03/1985")]:
            self.assertNotEqual(factory.fake_for(kind, value), value)


class DocxRoundTrip(unittest.TestCase):
    def _roundtrip(self, build):
        with tempfile.TemporaryDirectory() as folder:
            source, target = os.path.join(folder, "in.docx"), os.path.join(folder, "out.docx")
            document = Document()
            build(document)
            document.save(source)
            Redactor().redact_docx(source, target)
            return Document(target)

    def test_pii_split_across_runs_is_replaced_and_formatting_kept(self):
        def build(document):
            paragraph = document.add_paragraph()
            paragraph.add_run("Write to rohan.").bold = True
            paragraph.add_run("dey@gmail").italic = True
            paragraph.add_run(".com or call +91 98765 43210.")
        out = self._roundtrip(build)
        text = out.paragraphs[-1].text
        self.assertRegex(text, r"^Write to [\w.]+@example\.com or call \+91 [\d ]+\.$")
        self.assertNotIn("rohan", text)
        self.assertNotIn("98765", text)
        self.assertTrue(out.paragraphs[-1].runs[0].bold)

    def test_table_and_header_are_processed(self):
        def build(document):
            document.sections[0].header.paragraphs[0].text = "Contact: hr@corp.com"
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "SSN 123-45-6789"
        out = self._roundtrip(build)
        self.assertNotIn("123-45-6789", out.tables[0].cell(0, 0).text)
        self.assertNotIn("hr@corp.com", out.sections[0].header.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
